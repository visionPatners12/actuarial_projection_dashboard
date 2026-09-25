import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from model import BRANCHES, empty_direct_start, empty_reass_start
from exporter import export_optimized_forecast
from optimized_forecast import (
    ForecastInputError, calibrated_profile, run_optimized_forecast,
)


def fixture_inputs():
    history = pd.DataFrame([
        {"period": f"2026-{m:02d}", "branch": branch, "gwp_ytd": 100.0*m}
        for branch in BRANCHES for m in range(1, 13)
    ])
    direct = empty_direct_start()
    reass = empty_reass_start()
    direct["period"] = "2026-12"
    reass["period"] = "2026-12"
    direct["gwp_ytd"] = 1200.0
    direct["upr_close"] = 100.0
    direct["case_close_prior"] = 50.0
    direct["ibnr_close_prior"] = 10.0
    reass["ceded_premium_ytd"] = 240.0
    reass["ceded_upr_close"] = 20.0
    reass["recoverable_case_close_prior"] = 10.0
    reass["recoverable_ibnr_close_prior"] = 2.0
    rate = {
        "sp_current": 60.0, "sp_global": 80.0,
        "settlement_current": 50.0, "ibnr_share_current": 20.0,
        "recourse_current": 0.0, "recourse_prior": 0.0,
        "rec_direct": 20.0, "commission": 10.0,
        "cession": 20.0, "recovery_current": 20.0,
        "rec_reass": 20.0, "reass_commission": 10.0,
        "ifrs_ibnr_factor": 1.2, "ifrs_rec_factor": 1.4,
    }
    rates = pd.DataFrame([{"branch": branch, **rate} for branch in BRANCHES])
    return history, rates, direct, reass


def test_year_end_target_and_accounting():
    h, rates, direct, reass = fixture_inputs()
    targets = pd.DataFrame([{"branch": "Automobile", "gwp_end": 1500.0}])
    result = run_optimized_forecast(h, rates, direct, reass, targets,
                                    projection_year=2027)
    auto = result.direct[result.direct.branch == "Automobile"].sort_values("period")
    assert auto.iloc[-1].gwp_ytd == pytest.approx(1500.0)
    assert auto.iloc[0].upr_open == 100.0
    assert auto.upr_open.nunique() == 1
    assert np.all(np.diff(auto.gwp_ytd) >= -1e-8)
    assert np.allclose(
        auto.incurred_current_ytd,
        auto.paid_current_ytd-auto.recourse_current_ytd+
        auto.case_close_current+auto.ibnr_close_current-
        auto.case_open_current-auto.ibnr_open_current,
    )
    rr = result.reass[result.reass.branch == "Automobile"]
    assert np.allclose(rr.ceded_upr_close, rr.ceded_premium_ytd*0.2)
    assert np.allclose(rr.recovered_incurred_current_ytd,
                       auto.incurred_current_ytd*0.2)


def test_realized_months_and_point_lock():
    h, rates, direct, reass = fixture_inputs()
    baseline = run_optimized_forecast(h, rates, direct, reass,
                                      projection_year=2027)
    actual_d = pd.concat([direct, baseline.direct[
        baseline.direct.period.dt.month <= 2]], ignore_index=True)
    actual_r = pd.concat([reass, baseline.reass[
        baseline.reass.period.dt.month <= 2]], ignore_index=True)
    locks = pd.DataFrame([{
        "period": "2027-03", "branch": "Automobile",
        "field": "gwp_ytd", "mode": "cumul", "value": 350.0,
    }])
    targets = pd.DataFrame([{"branch": "Automobile", "gwp_end": 1300.0}])
    result = run_optimized_forecast(h, rates, actual_d, actual_r,
                                    annual_targets=targets, locks=locks,
                                    projection_year=2027)
    auto = result.direct[result.direct.branch == "Automobile"].sort_values("period")
    assert auto.iloc[1].gwp_ytd == pytest.approx(200.0)
    assert auto.iloc[2].gwp_ytd == pytest.approx(350.0)
    assert auto.iloc[-1].gwp_ytd == pytest.approx(1300.0)


def test_conflicting_target_fails():
    h, rates, direct, reass = fixture_inputs()
    locks = pd.DataFrame([{
        "period": "2027-03", "branch": "Automobile",
        "field": "gwp_ytd", "mode": "cumul", "value": 800.0,
    }])
    targets = pd.DataFrame([{"branch": "Automobile", "gwp_end": 500.0}])
    with pytest.raises(ForecastInputError):
        run_optimized_forecast(h, rates, direct, reass,
                               annual_targets=targets, locks=locks,
                               projection_year=2027)


def test_missing_rate_blocks_branch():
    h, rates, direct, reass = fixture_inputs()
    rates.loc[rates.branch == "Santé", "rec_reass"] = np.nan
    with pytest.raises(ForecastInputError, match="Santé.*rec_reass"):
        run_optimized_forecast(h, rates, direct, reass, projection_year=2027)


def test_portfolio_rebalance_preserves_priority_branches():
    h, rates, direct, reass = fixture_inputs()
    base = run_optimized_forecast(h, rates, direct, reass,
                                  projection_year=2027)
    target = pd.DataFrame([{
        "branch": "CONSOLIDATION", "sp_global": 85.0,
    }])
    result = run_optimized_forecast(h, rates, direct, reass,
                                    annual_targets=target,
                                    projection_year=2027)
    for branch in ("Automobile", "Santé"):
        before = base.direct[(base.direct.branch == branch) &
                             (base.direct.period.dt.month == 12)].iloc[0]
        after = result.direct[(result.direct.branch == branch) &
                              (result.direct.period.dt.month == 12)].iloc[0]
        assert after.incurred_current_ytd == pytest.approx(before.incurred_current_ytd)
    december = result.direct[result.direct.period.dt.month == 12]
    gross = (december.incurred_current_ytd+
             december.incurred_prior_ytd).sum()/december.earned_premium_ytd.sum()
    assert gross == pytest.approx(0.85)


def test_charge_lock_survives_portfolio_rebalance():
    h, rates, direct, reass = fixture_inputs()
    locks = pd.DataFrame([{
        "period": "2027-12", "branch": "Incendie",
        "field": "incurred_current_ytd", "mode": "cumul", "value": 700.0,
    }])
    target = pd.DataFrame([{"branch": "CONSOLIDATION", "sp_global": 85.0}])
    result = run_optimized_forecast(h, rates, direct, reass,
                                    annual_targets=target, locks=locks,
                                    projection_year=2027)
    dec = result.direct[result.direct.period.dt.month == 12]
    assert dec[dec.branch == "Incendie"].iloc[0].incurred_current_ytd == pytest.approx(700.0)
    gross = (dec.incurred_current_ytd+dec.incurred_prior_ytd).sum()/dec.earned_premium_ytd.sum()
    assert gross == pytest.approx(0.85)


def test_optimized_export_matches_result(tmp_path):
    h, rates, direct, reass = fixture_inputs()
    result = run_optimized_forecast(h, rates, direct, reass,
                                    projection_year=2027)
    path = tmp_path/"projection.xlsx"
    export_optimized_forecast(path, result, {"history_premiums": h})
    book = load_workbook(path, data_only=True)
    assert {"Direct Local", "Reass Local", "CPC SAZ Local",
            "Direct", "Reass", "CPC SAZ", "Contrôles"}.issubset(book.sheetnames)
    first = result.direct[result.direct.branch == "Automobile"].iloc[0]
    assert book["Direct Local"]["B7"].value == pytest.approx(first.gwp_ytd)
    assert book["Direct"]["B28"].value == pytest.approx(
        result.direct_ifrs[result.direct_ifrs.branch == "Automobile"].iloc[0].dac_open)


def test_one_remaining_month_and_ifrs_calibration():
    h, rates, direct, reass = fixture_inputs()
    base = run_optimized_forecast(h, rates, direct, reass, projection_year=2027)
    actual_d = pd.concat([
        direct, base.direct[base.direct.period.dt.month <= 11]], ignore_index=True)
    actual_r = pd.concat([
        reass, base.reass[base.reass.period.dt.month <= 11]], ignore_index=True)
    local = base.direct.merge(
        base.reass[["period", "branch", "ceded_upr_close"]],
        on=["period", "branch"])
    ifrs = base.direct_ifrs.merge(
        base.reass_ifrs[["period", "branch", "ceded_upr_close"]],
        on=["period", "branch"])
    rates["ifrs_ibnr_factor"] = np.nan
    rates["ifrs_rec_factor"] = np.nan
    result = run_optimized_forecast(
        h, rates, actual_d, actual_r, history_local=local,
        history_ifrs=ifrs, projection_year=2027)
    assert len(result.direct) == 96
    assert result.direct[result.direct.period.dt.month == 11].shape[0] == 8
    assert result.direct_ifrs[
        result.direct_ifrs.branch == "Automobile"].iloc[-1].ibnr_close_current > 0


def test_walk_forward_weights_use_available_years():
    h, _, _, _ = fixture_inputs()
    history = pd.concat([
        h.assign(period=h.period.str.replace("2026", "2024"),
                 gwp_ytd=h.gwp_ytd*0.7),
        h.assign(period=h.period.str.replace("2026", "2025"),
                 gwp_ytd=h.gwp_ytd*0.9),
        h,
    ], ignore_index=True)
    projected, weights, mae = calibrated_profile(
        history, "Automobile", 2026)
    assert len(projected) == 12
    assert sum(weights.values()) == pytest.approx(1.0)
    assert all(w >= 0 for w in weights.values())
    assert mae is not None


def test_monthly_ifrs_target_changes_only_ifrs():
    h, rates, direct, reass = fixture_inputs()
    baseline = run_optimized_forecast(h, rates, direct, reass,
                                      projection_year=2027)
    march = baseline.direct_ifrs[
        (baseline.direct_ifrs.branch == "Automobile") &
        (baseline.direct_ifrs.period.dt.month == 3)].iloc[0]
    target_sp = march.sp_exercice*100+5
    exceptions = pd.DataFrame([{
        "period": "2027-03", "branch": "Automobile",
        "sp_ifrs_current": target_sp,
    }])
    result = run_optimized_forecast(
        h, rates, direct, reass, rate_exceptions=exceptions,
        projection_year=2027)
    local_march = result.direct[
        (result.direct.branch == "Automobile") &
        (result.direct.period.dt.month == 3)].iloc[0]
    ifrs_march = result.direct_ifrs[
        (result.direct_ifrs.branch == "Automobile") &
        (result.direct_ifrs.period.dt.month == 3)].iloc[0]
    assert local_march.incurred_current_ytd == pytest.approx(
        baseline.direct[
            (baseline.direct.branch == "Automobile") &
            (baseline.direct.period.dt.month == 3)].iloc[0].incurred_current_ytd)
    assert ifrs_march.sp_exercice*100 == pytest.approx(target_sp)


def test_exception_cannot_change_realized_month():
    h, rates, direct, reass = fixture_inputs()
    baseline = run_optimized_forecast(h, rates, direct, reass,
                                      projection_year=2027)
    actual_d = pd.concat([direct, baseline.direct[
        baseline.direct.period.dt.month == 1]], ignore_index=True)
    actual_r = pd.concat([reass, baseline.reass[
        baseline.reass.period.dt.month == 1]], ignore_index=True)
    exceptions = pd.DataFrame([{
        "period": "2027-01", "branch": "Automobile", "sp_current": 90.0,
    }])
    with pytest.raises(ForecastInputError, match="mois réalisé"):
        run_optimized_forecast(h, rates, actual_d, actual_r,
                               rate_exceptions=exceptions, projection_year=2027)
