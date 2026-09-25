from __future__ import annotations
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Tuple, List
import pandas as pd
import numpy as np
from openpyxl import load_workbook

from model import BRANCHES, empty_direct_start, empty_reass_start

DIRECT_ROWS = {
    "gwp_ytd": 7, "pap_open": 8, "pap_close": 9, "pane_open": 10, "pane_close": 11,
    "revenue": 12, "paid_current_ytd": 14, "paid_prior_ytd": 15, "commission_ytd": 18,
    "upr_open": 20, "upr_close": 21, "ibnr_open_total": 36, "ibnr_close_total": 37,
    "ibnr_be_open_total": 58, "ibnr_be_close_current": 59, "ibnr_be_close_prior": 60,
    "case_close_current": 45, "case_close_prior": 46, "case_open_total": 65,
    "case_close_current_detail": 66, "case_close_prior_detail": 67, "recourse_current_ytd": 47,
    "recourse_prior_ytd": 48,
}

REASS_ROWS = {
    "ceded_premium_ytd": 7, "recovered_paid_current_ytd": 14, "recovered_paid_prior_ytd": 15,
    "reass_commission_ytd": 18, "ceded_upr_open": 20, "ceded_upr_close": 21,
    "recoverable_ibnr_open_total": 36, "recoverable_ibnr_close_total": 37,
    "recoverable_ibnr_be_open_total": 55, "recoverable_ibnr_be_close_current": 56, "recoverable_ibnr_be_close_prior": 57,
    "recoverable_case_close_current": 44, "recoverable_case_close_prior": 45,
    "recoverable_case_open_total": 62, "recoverable_case_close_current_detail": 63, "recoverable_case_close_prior_detail": 64,
}


def _num(v):
    if v is None:
        return 0.0
    try:
        if pd.isna(v): return 0.0
    except Exception:
        pass
    try: return float(v)
    except Exception: return 0.0


def _block_starts(ws) -> List[int]:
    starts=[]
    for c in range(1, ws.max_column+1):
        if ws.cell(6,c).value == "Automobile":
            starts.append(c)
    return starts


def _infer_year(ws, starts: List[int]) -> int:
    years=[]
    for c in starts:
        for j in range(c, min(ws.max_column,c+11)+1):
            v=ws.cell(4,j).value
            if isinstance(v, datetime): years.append(v.year)
    return Counter(years).most_common(1)[0][0] if years else datetime.today().year


def _extract_sheet(ws_values, starts: List[int], year: int, rows: dict, kind: str) -> pd.DataFrame:
    out=[]
    for bi,c0 in enumerate(starts):
        month=bi+1
        if month>12: break
        period=pd.Timestamp(year=year,month=month,day=1)+pd.offsets.MonthEnd(0)
        branch_names=[ws_values.cell(6,c0+k).value for k in range(8)]
        for k,b in enumerate(branch_names):
            if b not in BRANCHES:
                continue
            c=c0+k
            rec={"period":period,"branch":b}
            for key,r in rows.items():
                rec[key]=_num(ws_values.cell(r,c).value)
            out.append(rec)
    df=pd.DataFrame(out)
    if df.empty: return df
    # Reconstruct missing split while respecting the model rule that openings are
    # annual balances: for a given branch/year they remain fixed on every monthly
    # block. The legacy sheet only exposes aggregate IBNR opening and no explicit
    # SAP opening split, so we use the first period of the year as the annual prior
    # opening and keep current-year opening at zero.
    for (b, year), idx in df.groupby(["branch", df["period"].dt.year]).groups.items():
        inds=sorted(list(idx), key=lambda i: df.loc[i,"period"])
        first_i=inds[0]
        if kind=="direct":
            # Prefer the detailed BE/SAP section when it is populated. Annual
            # opening balances belong to prior years at 1 January; current-year
            # opening is therefore zero. Fall back to the older aggregate rows
            # only when the detailed section is absent/empty.
            detailed_case_open=_num(df.loc[first_i,"case_open_total"])
            fallback_case_open=_num(df.loc[first_i,"case_close_prior"])
            annual_case_open_current=0.0
            annual_case_open_prior=max(0.0, detailed_case_open if abs(detailed_case_open)>1e-9 else fallback_case_open)
            detailed_ibnr_open=_num(df.loc[first_i,"ibnr_be_open_total"])
            fallback_ibnr_open=_num(df.loc[first_i,"ibnr_open_total"])
            annual_ibnr_open_current=0.0
            annual_ibnr_open_prior=max(0.0, detailed_ibnr_open if abs(detailed_ibnr_open)>1e-9 else fallback_ibnr_open)
            for i in inds:
                df.loc[i,"case_open_current"]=annual_case_open_current
                df.loc[i,"case_open_prior"]=annual_case_open_prior
                df.loc[i,"ibnr_open_current"]=annual_ibnr_open_current
                df.loc[i,"ibnr_open_prior"]=annual_ibnr_open_prior
                cc=_num(df.loc[i,"case_close_current_detail"]); cp=_num(df.loc[i,"case_close_prior_detail"])
                if abs(cc)+abs(cp) <= 1e-9:
                    cc=_num(df.loc[i,"case_close_current"]); cp=_num(df.loc[i,"case_close_prior"])
                df.loc[i,"case_close_current"]=max(0.0,cc)
                df.loc[i,"case_close_prior"]=max(0.0,cp)
                ic=_num(df.loc[i,"ibnr_be_close_current"]); ip=_num(df.loc[i,"ibnr_be_close_prior"])
                if abs(ic)+abs(ip) <= 1e-9:
                    ic=0.0; ip=_num(df.loc[i,"ibnr_close_total"])
                df.loc[i,"ibnr_close_current"]=max(0.0,ic)
                df.loc[i,"ibnr_close_prior"]=max(0.0,ip)
        else:
            detailed_case_open=_num(df.loc[first_i,"recoverable_case_open_total"])
            fallback_case_open=_num(df.loc[first_i,"recoverable_case_close_prior"])
            annual_case_open_current=0.0
            annual_case_open_prior=max(0.0, detailed_case_open if abs(detailed_case_open)>1e-9 else fallback_case_open)
            detailed_ibnr_open=_num(df.loc[first_i,"recoverable_ibnr_be_open_total"])
            fallback_ibnr_open=_num(df.loc[first_i,"recoverable_ibnr_open_total"])
            annual_ibnr_open_current=0.0
            annual_ibnr_open_prior=max(0.0, detailed_ibnr_open if abs(detailed_ibnr_open)>1e-9 else fallback_ibnr_open)
            for i in inds:
                df.loc[i,"recoverable_case_open_current"]=annual_case_open_current
                df.loc[i,"recoverable_case_open_prior"]=annual_case_open_prior
                df.loc[i,"recoverable_ibnr_open_current"]=annual_ibnr_open_current
                df.loc[i,"recoverable_ibnr_open_prior"]=annual_ibnr_open_prior
                cc=_num(df.loc[i,"recoverable_case_close_current_detail"]); cp=_num(df.loc[i,"recoverable_case_close_prior_detail"])
                if abs(cc)+abs(cp) <= 1e-9:
                    cc=_num(df.loc[i,"recoverable_case_close_current"]); cp=_num(df.loc[i,"recoverable_case_close_prior"])
                df.loc[i,"recoverable_case_close_current"]=max(0.0,cc)
                df.loc[i,"recoverable_case_close_prior"]=max(0.0,cp)
                ic=_num(df.loc[i,"recoverable_ibnr_be_close_current"]); ip=_num(df.loc[i,"recoverable_ibnr_be_close_prior"])
                if abs(ic)+abs(ip) <= 1e-9:
                    ic=0.0; ip=_num(df.loc[i,"recoverable_ibnr_close_total"])
                df.loc[i,"recoverable_ibnr_close_current"]=max(0.0,ic)
                df.loc[i,"recoverable_ibnr_close_prior"]=max(0.0,ip)
    return df



def _history_coherence_messages(df: pd.DataFrame, kind: str) -> List[str]:
    msgs=[]
    if df is None or df.empty:
        return msgs
    d=df.copy(); d["period"]=pd.to_datetime(d["period"])
    prefix="Direct" if kind=="direct" else "Réassurance"
    if kind=="direct":
        for f in ["upr_open","upr_close"]:
            vals=pd.to_numeric(d.get(f),errors="coerce")
            n=int((vals<0).sum())
            if n:
                msgs.append(f"⚠️ {prefix}: {n} valeur(s) REC négative(s) détectée(s) dans l'historique ({f}). Elles ne seront pas reproduites par le moteur.")
        if all(c in d.columns for c in ["gwp_ytd","upr_open","upr_close"]):
            d["earned_check"]=pd.to_numeric(d["gwp_ytd"],errors="coerce").fillna(0)+pd.to_numeric(d["upr_open"],errors="coerce").fillna(0)-pd.to_numeric(d["upr_close"],errors="coerce").fillna(0)
            ndec=0
            for (_, _),g in d.groupby(["branch",d["period"].dt.year]):
                vals=g.sort_values("period")["earned_check"].to_numpy(float)
                ndec += int((np.diff(vals)<-1.0).sum()) if len(vals)>1 else 0
            if ndec:
                msgs.append(f"⚠️ {prefix}: {ndec} baisse(s) de prime acquise cumulée détectée(s) dans l'historique. Le nouveau moteur impose une trajectoire non décroissante.")
        for f in ["pap_open","pap_close","pane_open","pane_close"]:
            if f in d.columns:
                varied=0
                for (_, _),g in d.groupby(["branch",d["period"].dt.year]):
                    vals=pd.to_numeric(g[f],errors="coerce").dropna().to_numpy(float)
                    if len(vals)>1 and np.nanmax(vals)-np.nanmin(vals)>1.0:
                        varied+=1
                if varied:
                    msgs.append(f"⚠️ {prefix}: {f} varie dans {varied} branche/exercice(s). Dans le nouveau moteur PAP/PANE restent fixes sur l'exercice.")
    else:
        for f in ["ceded_upr_open","ceded_upr_close"]:
            vals=pd.to_numeric(d.get(f),errors="coerce")
            n=int((vals<0).sum())
            if n:
                msgs.append(f"⚠️ {prefix}: {n} valeur(s) REC cédée négative(s) détectée(s) ({f}). Elles ne seront pas reproduites.")
        if all(c in d.columns for c in ["ceded_premium_ytd","ceded_upr_open","ceded_upr_close"]):
            d["earned_check"]=pd.to_numeric(d["ceded_premium_ytd"],errors="coerce").fillna(0)+pd.to_numeric(d["ceded_upr_open"],errors="coerce").fillna(0)-pd.to_numeric(d["ceded_upr_close"],errors="coerce").fillna(0)
            ndec=0
            for (_, _),g in d.groupby(["branch",d["period"].dt.year]):
                vals=g.sort_values("period")["earned_check"].to_numpy(float)
                ndec += int((np.diff(vals)<-1.0).sum()) if len(vals)>1 else 0
            if ndec:
                msgs.append(f"⚠️ {prefix}: {ndec} baisse(s) de prime acquise cédée cumulée détectée(s). Le nouveau moteur impose une trajectoire non décroissante.")
    return msgs

def parse_legacy_workbook(path: str):
    p=Path(path)
    wbv=load_workbook(p, data_only=True, read_only=False)
    messages=[]
    direct_hist=pd.DataFrame(); reass_hist=pd.DataFrame()
    direct_year=reass_year=None

    if "Direct Local" in wbv.sheetnames:
        ws=wbv["Direct Local"]
        starts=_block_starts(ws)
        direct_year=_infer_year(ws,starts)
        direct_hist=_extract_sheet(ws,starts,direct_year,DIRECT_ROWS,"direct")
        messages.append(f"Direct Local: {len(starts)} blocs détectés, année de référence {direct_year}.")
    else:
        messages.append("Feuille 'Direct Local' absente.")

    if "Reass Local" in wbv.sheetnames:
        ws=wbv["Reass Local"]
        starts=_block_starts(ws)
        reass_year=_infer_year(ws,starts)
        reass_hist=_extract_sheet(ws,starts,reass_year,REASS_ROWS,"reass")
        messages.append(f"Reass Local: {len(starts)} blocs détectés, année de référence {reass_year}.")
    else:
        messages.append("Feuille 'Reass Local' absente.")

    if direct_year is not None and reass_year is not None and direct_year != reass_year:
        messages.append(f"⚠️ Incohérence de calendrier détectée: Direct={direct_year}, Réassurance={reass_year}. Le moteur ne fusionne pas ces dates silencieusement.")

    messages.extend(_history_coherence_messages(direct_hist, "direct"))
    messages.extend(_history_coherence_messages(reass_hist, "reass"))

    sd=empty_direct_start(); sr=empty_reass_start()
    if not direct_hist.empty:
        last_period=direct_hist["period"].max()
        last=direct_hist[direct_hist["period"]==last_period]
        for i,b in enumerate(sd["branch"]):
            h=last[last["branch"]==b]
            if h.empty: continue
            r=h.iloc[0]
            for c in sd.columns:
                if c!="branch" and c in r.index: sd.loc[i,c]=_num(r[c])
        start_period=str(pd.Period(last_period,freq="M"))
    elif not reass_hist.empty:
        start_period=str(pd.Period(reass_hist["period"].max(),freq="M"))
    else:
        start_period=str(pd.Period(datetime.today(),freq="M"))

    if not reass_hist.empty and not (direct_year is not None and reass_year is not None and direct_year != reass_year):
        last_period_r=reass_hist["period"].max()
        last=reass_hist[reass_hist["period"]==last_period_r]
        for i,b in enumerate(sr["branch"]):
            h=last[last["branch"]==b]
            if h.empty: continue
            r=h.iloc[0]
            for c in sr.columns:
                if c!="branch" and c in r.index: sr.loc[i,c]=_num(r[c])
    elif not reass_hist.empty:
        messages.append("⚠️ Le bloc de départ Réassurance n'est pas prérempli car son calendrier ne correspond pas à Direct. Les taux de cession/récupération seront estimés depuis l'historique disponible ou saisis dans Hyp Reass.")

    messages.append("Le nouveau moteur utilise des % transparents : progression mensuelle pour les flux, taux sur primes pour commissions/REC, taux de cession et récupération pour la réassurance.")
    messages.append("Les ouvertures restent fixes dans l'exercice; PAP/PANE restent fixes; la prime acquise cumulée ne peut pas diminuer.")
    return direct_hist, reass_hist, sd, sr, start_period, messages
