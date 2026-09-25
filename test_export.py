import tempfile
import pandas as pd
from openpyxl import load_workbook
from model import *
from exporter import export_projection


def test_n(n):
    sd=empty_direct_start(); sr=empty_reass_start();
    # seed all branches to avoid completely empty output
    sd['gwp_ytd']=1_000_000_000; sd['upr_open']=200_000_000; sd['upr_close']=180_000_000
    sd['paid_current_ytd']=100_000_000; sd['case_close_current']=50_000_000; sd['ibnr_close_current']=20_000_000
    start='2026-12'; hd=pd.DataFrame(); hr=pd.DataFrame()
    h=build_hyp_direct(hd,start,n); rh=build_hyp_reass(hd,hr,start,n)
    cd=empty_cibles_direct(start,n); cr=empty_cibles_reass(start,n)
    d,r,s,diag,apd,apr,*_=run_projection(sd,sr,h,rh,cd,cr,start,n,hd,hr)
    out=tempfile.mktemp(suffix='.xlsx')
    export_projection(out,{'cibles_direct':cd,'cibles_reass':cr},d,r,s,diag,apd,apr)
    wb=load_workbook(out,data_only=True)
    assert wb.sheetnames==['Hyp Direct','Hyp Reass','Cibles','Direct Local','Reass Local','CPC SAZ']
    # Detect month blocks by Automobile header.
    for sheet in ['Direct Local','Reass Local']:
        ws=wb[sheet]; starts=[c for c in range(1,ws.max_column+1) if ws.cell(6,c).value=='Automobile']
        assert len(starts)==n,(sheet,n,len(starts))
    # CPC monthly block starts contain a month label in row 1 and CONSOLIDATION in row2.
    ws=wb['CPC SAZ']; starts=[c for c in range(1,ws.max_column+1) if ws.cell(2,c).value=='CONSOLIDATION ']
    assert len(starts)==n,(n,len(starts))
    # Percentage block: Jan blank, Feb populated if present.
    wd=wb['Direct Local']; wr=wb['Reass Local']
    if n>=2:
        dstarts=[c for c in range(1,wd.max_column+1) if wd.cell(6,c).value=='Automobile']
        rstarts=[c for c in range(1,wr.max_column+1) if wr.cell(6,c).value=='Automobile']
        # If second period is February (start Dec -> Jan, Feb), it must show %.
        assert wd.cell(134,dstarts[1]).value is not None
        assert wr.cell(90,rstarts[1]).value is not None
    return out

if __name__=='__main__':
    for n in [1,4,12,18,24,36,48]:
        f=test_n(n); print('PASS export',n,f)
