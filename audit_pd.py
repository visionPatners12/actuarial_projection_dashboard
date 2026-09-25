"""Create a row-level, reproducible inventory of the reference workbook.

This inventory records observed formulas; it never treats them as model rules.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(__file__).with_name("pd.xlsx")
OUTPUT = Path(__file__).with_name("AUDIT_PD.csv")
ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#NUM!", "#N/A", "#NAME?"}


def classify(label):
    text = label.casefold().strip()
    if any(x in text for x in ("ouverture", "ouv")):
        return "FIXE", "Stock"
    if any(x in text for x in ("s/p", "ratio", "taux", "multiplicateur", " k ")):
        return "TAUX × BASE", "Ratio"
    if any(x in text for x in ("variation", "var ", "chiffre d'affaire",
                              "primes acquises", "charge globale",
                              "charges de sinistres", "résultat",
                              "consolidation", "net")):
        return "IDENTITÉ COMPTABLE", "Calcul"
    if any(x in text for x in ("ibnr", "sap", "provision")):
        return "VARIABLE ACTUARIELLE", "Stock ou mouvement"
    if "clôture" in text or "cloture" in text or "clo" in text:
        return "PROJECTION TEMPORELLE", "Stock"
    if any(x in text for x in ("prime", "sinistre", "recours", "commission",
                              "dac", "rec ")):
        return "PROJECTION TEMPORELLE", "Flux ou stock selon la ligne"
    return "À QUALIFIER", "En-tête ou autre"


def audit(source=SOURCE, output=OUTPUT):
    wb = load_workbook(source, data_only=False, read_only=False)
    rows = []
    for ws in wb:
        label_col = 2 if ws.title.startswith("CPC SAZ") else 1
        for row in ws.iter_rows():
            label_cell = row[label_col-1]
            label = label_cell.value
            formulas = [
                c for c in row if isinstance(c.value, str) and
                c.value.startswith("=")
            ]
            errors = [c.coordinate for c in row
                      if isinstance(c.value, str) and c.value in ERRORS]
            if not isinstance(label, str) or not label.strip() or label.startswith("="):
                if not formulas and not errors:
                    continue
                label = "(sans libellé)"
            category, nature = classify(label)
            first = formulas[0] if formulas else None
            sample = str(first.value)[:220] if first else ""
            references = sorted(set(re.findall(
                r"'([^']+)'!|([A-Za-z_][A-Za-z0-9_ ]*)!",
                sample)))
            refs = ", ".join(a or b for a, b in references)
            rows.append({
                "onglet": ws.title, "ligne": label_cell.row,
                "libelle": label.strip(), "categorie_proposee": category,
                "nature_proposee": nature,
                "premiere_formule_cellule": first.coordinate if first else "",
                "formule_observee": sample, "onglets_references": refs,
                "erreurs_litterales": ", ".join(errors[:12]),
                "nombre_formules": len(formulas),
                "statut": "À valider métier ; formule non reprise automatiquement",
            })
    with open(output, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return rows


if __name__ == "__main__":
    result = audit()
    print(f"{len(result)} lignes auditées dans {OUTPUT}")
