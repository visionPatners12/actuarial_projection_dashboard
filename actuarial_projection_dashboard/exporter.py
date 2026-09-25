from __future__ import annotations
from pathlib import Path
import tempfile
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

RATE_HINTS = ("rate", "sp_", "loss_ratio", "cession", "recovery")


def _write_df(ws, df: pd.DataFrame, input_sheet=False):
    if df is None: df = pd.DataFrame()
    data = df.copy()
    for c in data.columns:
        if pd.api.types.is_datetime64_any_dtype(data[c]):
            data[c] = data[c].dt.strftime("%Y-%m-%d")
    ws.append(list(data.columns))
    for row in data.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])
    dark = PatternFill("solid", fgColor="17365D")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="D9E1F2")
    for cell in ws[1]:
        cell.fill = dark; cell.font = Font(color="FFFFFF", bold=True); cell.alignment=Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = Border(bottom=thin)
            if input_sheet and isinstance(cell.value,(int,float)):
                cell.font = Font(color="0000FF")
                cell.fill = input_fill
            elif isinstance(cell.value,(int,float)):
                cell.font = Font(color="000000")
    for idx,col in enumerate(data.columns,1):
        width=max(12,min(34,max([len(str(col))]+[len(str(v)) for v in data[col].head(50).fillna("")])+2))
        ws.column_dimensions[get_column_letter(idx)].width=width
        if any(h in str(col).lower() for h in RATE_HINTS):
            for cell in ws[get_column_letter(idx)][1:]: cell.number_format='0.0%'
        elif pd.api.types.is_numeric_dtype(data[col]):
            for cell in ws[get_column_letter(idx)][1:]: cell.number_format='#,##0;[Red](#,##0);-'
    ws.freeze_panes="A2"
    ws.sheet_view.showGridLines=False


def export_projection(path, inputs: dict, direct, reass, summary, diagnostics, params, direct_block, reass_block):
    wb=Workbook(); wb.remove(wb.active)
    sheets=[
        ("Start Direct", inputs.get("start_direct"), True),
        ("Start Reass", inputs.get("start_reass"), True),
        ("Targets Direct", inputs.get("target_direct"), True),
        ("Targets Reass", inputs.get("target_reass"), True),
        ("Hypotheses", inputs.get("assumptions"), True),
        ("Projection Direct", direct, False),
        ("Projection Reass", reass, False),
        ("Bloc Direct", direct_block, False),
        ("Bloc Reass", reass_block, False),
        ("Synthese", summary, False),
        ("Parametres calibres", params, False),
        ("Diagnostics", diagnostics, False),
    ]
    for name,df,inp in sheets:
        ws=wb.create_sheet(name[:31]); _write_df(ws, pd.DataFrame(df) if df is not None else pd.DataFrame(), inp)
    wb.save(path)
    return path
