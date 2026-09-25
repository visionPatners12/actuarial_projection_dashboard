from __future__ import annotations
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Tuple, List
import pandas as pd
import numpy as np
from openpyxl import load_workbook

from model import BRANCHES, empty_direct_start, empty_reass_start, default_assumptions, estimate_profiles

DIRECT_ROWS = {
    "gwp_ytd": 7, "pap_open": 8, "pap_close": 9, "pane_open": 10, "pane_close": 11,
    "revenue": 12, "paid_current_ytd": 14, "paid_prior_ytd": 15, "commission_ytd": 18,
    "upr_open": 20, "upr_close": 21, "ibnr_open_total": 36, "ibnr_close_total": 37,
    "case_close_current": 45, "case_close_prior": 46, "recourse_current_ytd": 47,
    "recourse_prior_ytd": 48,
}

REASS_ROWS = {
    "ceded_premium_ytd": 7, "recovered_paid_current_ytd": 14, "recovered_paid_prior_ytd": 15,
    "reass_commission_ytd": 18, "ceded_upr_open": 20, "ceded_upr_close": 21,
    "recoverable_ibnr_open_total": 36, "recoverable_ibnr_close_total": 37,
    "recoverable_case_close_current": 44, "recoverable_case_close_prior": 45,
}


def _num(v):
    if v is None:
        return 0.0
    try:
        if pd.isna(v): return 0.0
    except Exception:
        pass
    try: return float(v)
    except Exception: return 0.0


def _block_starts(ws) -> List[int]:
    starts=[]
    for c in range(1, ws.max_column+1):
        if ws.cell(6,c).value == "Automobile":
            starts.append(c)
    return starts


def _infer_year(ws, starts: List[int]) -> int:
    years=[]
    for c in starts:
        for j in range(c, min(ws.max_column,c+11)+1):
            v=ws.cell(4,j).value
            if isinstance(v, datetime): years.append(v.year)
    return Counter(years).most_common(1)[0][0] if years else datetime.today().year


def _extract_sheet(ws_values, starts: List[int], year: int, rows: dict, kind: str) -> pd.DataFrame:
    out=[]
    for bi,c0 in enumerate(starts):
        month=bi+1
        if month>12: break
        period=pd.Timestamp(year=year,month=month,day=1)+pd.offsets.MonthEnd(0)
        branch_names=[ws_values.cell(6,c0+k).value for k in range(8)]
        for k,b in enumerate(branch_names):
            if b not in BRANCHES:
                continue
            c=c0+k
            rec={"period":period,"branch":b}
            for key,r in rows.items():
                rec[key]=_num(ws_values.cell(r,c).value)
            out.append(rec)
    df=pd.DataFrame(out)
    if df.empty: return df
    # Reconstruct opening split from prior closing. File has many aggregate openings only.
    for b, idx in df.groupby("branch").groups.items():
        inds=list(idx)
        inds=sorted(inds,key=lambda i: df.loc[i,"period"])
        if kind=="direct":
            for pos,i in enumerate(inds):
                if pos==0:
                    df.loc[i,"case_open_current"]=0.0
                    df.loc[i,"case_open_prior"]=max(0.0,_num(df.loc[i,"case_close_prior"]))
                    df.loc[i,"ibnr_open_current"]=0.0
                    df.loc[i,"ibnr_open_prior"]=_num(df.loc[i,"ibnr_open_total"])
                else:
                    p=inds[pos-1]
                    df.loc[i,"case_open_current"]=_num(df.loc[p,"case_close_current"])
                    df.loc[i,"case_open_prior"]=_num(df.loc[p,"case_close_prior"])
                    df.loc[i,"ibnr_open_current"]=_num(df.loc[p,"ibnr_close_current"])
                    df.loc[i,"ibnr_open_prior"]=_num(df.loc[p,"ibnr_close_prior"])
                # The legacy workbook exposes aggregate CIMA IBNR. Preserve it as prior when split is unavailable.
                df.loc[i,"ibnr_close_current"]=0.0
                df.loc[i,"ibnr_close_prior"]=_num(df.loc[i,"ibnr_close_total"])
        else:
            for pos,i in enumerate(inds):
                if pos==0:
                    df.loc[i,"recoverable_case_open_current"]=0.0
                    df.loc[i,"recoverable_case_open_prior"]=_num(df.loc[i,"recoverable_case_close_prior"])
                    df.loc[i,"recoverable_ibnr_open_current"]=0.0
                    df.loc[i,"recoverable_ibnr_open_prior"]=_num(df.loc[i,"recoverable_ibnr_open_total"])
                else:
                    p=inds[pos-1]
                    df.loc[i,"recoverable_case_open_current"]=_num(df.loc[p,"recoverable_case_close_current"])
                    df.loc[i,"recoverable_case_open_prior"]=_num(df.loc[p,"recoverable_case_close_prior"])
                    df.loc[i,"recoverable_ibnr_open_current"]=_num(df.loc[p,"recoverable_ibnr_close_current"])
                    df.loc[i,"recoverable_ibnr_open_prior"]=_num(df.loc[p,"recoverable_ibnr_close_prior"])
                df.loc[i,"recoverable_ibnr_close_current"]=0.0
                df.loc[i,"recoverable_ibnr_close_prior"]=_num(df.loc[i,"recoverable_ibnr_close_total"])
    return df


def parse_legacy_workbook(path: str):
    p=Path(path)
    wbv=load_workbook(p, data_only=True, read_only=False)
    messages=[]
    direct_hist=pd.DataFrame(); reass_hist=pd.DataFrame()
    direct_year=reass_year=None

    if "Direct Local" in wbv.sheetnames:
        ws=wbv["Direct Local"]
        starts=_block_starts(ws)
        direct_year=_infer_year(ws,starts)
        direct_hist=_extract_sheet(ws,starts,direct_year,DIRECT_ROWS,"direct")
        messages.append(f"Direct Local: {len(starts)} blocs détectés, année de référence {direct_year}.")
    else:
        messages.append("Feuille 'Direct Local' absente.")

    if "Reass Local" in wbv.sheetnames:
        ws=wbv["Reass Local"]
        starts=_block_starts(ws)
        reass_year=_infer_year(ws,starts)
        reass_hist=_extract_sheet(ws,starts,reass_year,REASS_ROWS,"reass")
        messages.append(f"Reass Local: {len(starts)} blocs détectés, année de référence {reass_year}.")
    else:
        messages.append("Feuille 'Reass Local' absente.")

    if direct_year is not None and reass_year is not None and direct_year != reass_year:
        messages.append(f"⚠️ Incohérence de calendrier détectée: Direct={direct_year}, Réassurance={reass_year}. Le moteur ne fusionne pas ces dates silencieusement.")

    sd=empty_direct_start(); sr=empty_reass_start(); ah=default_assumptions()
    if not direct_hist.empty:
        last_period=direct_hist["period"].max()
        last=direct_hist[direct_hist["period"]==last_period]
        for i,b in enumerate(sd["branch"]):
            h=last[last["branch"]==b]
            if h.empty: continue
            r=h.iloc[0]
            for c in sd.columns:
                if c!="branch" and c in r.index: sd.loc[i,c]=_num(r[c])
        start_period=str(pd.Period(last_period,freq="M"))
    elif not reass_hist.empty:
        start_period=str(pd.Period(reass_hist["period"].max(),freq="M"))
    else:
        start_period=str(pd.Period(datetime.today(),freq="M"))

    if not reass_hist.empty and not (direct_year is not None and reass_year is not None and direct_year != reass_year):
        last_period_r=reass_hist["period"].max()
        last=reass_hist[reass_hist["period"]==last_period_r]
        for i,b in enumerate(sr["branch"]):
            h=last[last["branch"]==b]
            if h.empty: continue
            r=h.iloc[0]
            for c in sr.columns:
                if c!="branch" and c in r.index: sr.loc[i,c]=_num(r[c])
    elif not reass_hist.empty:
        messages.append("⚠️ Le bloc de départ Réassurance n'est pas prérempli car son calendrier ne correspond pas à Direct. Saisis/corrige l'ouverture Réass avant calcul.")

    # Surface historical estimates in the assumptions table; the user remains free to change mode/value.
    profiles=estimate_profiles(direct_hist,reass_hist,BRANCHES)
    for i,b in enumerate(ah["branch"]):
        prof=profiles.get(b,{})
        for key in ["sp_exercice","sp_global","commission_rate","annual_premium_growth","payment_rate_current","payment_rate_prior","ibnr_share_current","ibnr_share_prior","premium_cession_rate","claim_recovery_current","claim_recovery_prior","reass_commission_rate"]:
            if key in prof: ah.loc[i,key]=prof[key]
        h=direct_hist[direct_hist["branch"]==b] if not direct_hist.empty else pd.DataFrame()
        if not h.empty:
            lr=h.sort_values("period").iloc[-1]
            gwp=_num(lr.get("gwp_ytd")); upr=_num(lr.get("upr_close"))
            if gwp>0: ah.loc[i,"rec_target_rate"]=max(0.0,upr/gwp)

    messages.append("Les ouvertures du moteur seront prises du dernier état de clôture et resteront verrouillées pendant l'optimisation.")
    return direct_hist, reass_hist, sd, sr, ah, start_period, messages
